"""
Excel Manager Module.
Persists extracted lead intelligence into a styled Excel workbook (`.xlsx`)
with thread safety for concurrent executions.
"""

from typing import Dict, List, Any, Optional
from pathlib import Path
import threading
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import config


class ExcelManager:
    """
    Manages reading, creating, and appending extracted lead data into an Excel spreadsheet.
    Thread-safe for concurrent processing.
    """

    HEADERS = config.TARGET_SCHEMA

    def __init__(self, file_path: Optional[str] = None):
        """
        Initialize the Excel Manager.

        :param file_path: Path to the target .xlsx file. Defaults to `config.DEFAULT_EXCEL_PATH`.
        """
        self.file_path = Path(file_path or config.DEFAULT_EXCEL_PATH)
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self) -> None:
        """Creates the Excel workbook with formatted headers if it doesn't exist."""
        with self._lock:
            if not self.file_path.exists():
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Leads"

                # Header Styles
                header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Write headers
                ws.append(self.HEADERS)
                for col_idx in range(1, len(self.HEADERS) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment

                # Set initial column widths
                for col_idx, header in enumerate(self.HEADERS, 1):
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = max(len(header) + 6, 18)

                wb.save(str(self.file_path))
                wb.close()

    def append_lead(self, lead_data: Dict[str, Any]) -> None:
        """
        Append a single lead record to the Excel file.

        :param lead_data: Dictionary containing keys matching `config.TARGET_SCHEMA`.
        """
        self.append_leads_batch([lead_data])

    def append_leads_batch(self, leads: List[Dict[str, Any]]) -> None:
        """
        Append multiple lead records to the Excel file with auto-adjusting column widths.
        Includes retry logic for Windows file locks.

        :param leads: List of lead dictionaries.
        """
        if not leads:
            return

        with self._lock:
            max_retries = 4
            retry_delay = 0.5

            for attempt in range(max_retries):
                wb = None
                try:
                    wb = openpyxl.load_workbook(str(self.file_path))
                    ws = wb.active

                    thin_border = Border(
                        left=Side(style="thin", color="D9D9D9"),
                        right=Side(style="thin", color="D9D9D9"),
                        top=Side(style="thin", color="D9D9D9"),
                        bottom=Side(style="thin", color="D9D9D9")
                    )
                    body_font = Font(name="Calibri", size=10)
                    body_alignment = Alignment(vertical="center", wrap_text=True)

                    for lead in leads:
                        row_values = [lead.get(header, "") for header in self.HEADERS]
                        ws.append(row_values)
                        current_row = ws.max_row

                        for col_idx in range(1, len(self.HEADERS) + 1):
                            cell = ws.cell(row=current_row, column=col_idx)
                            cell.font = body_font
                            cell.alignment = body_alignment
                            cell.border = thin_border

                    # Auto-fit column widths
                    for col in ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        for cell in col:
                            if cell.value is not None:
                                lines = str(cell.value).split("\n")
                                line_lens = [len(l) for l in lines]
                                max_len = max(max_len, max(line_lens))
                        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 16)

                    wb.save(str(self.file_path))
                    return

                except PermissionError as pe:
                    if attempt < max_retries - 1:
                        import time
                        time.sleep(retry_delay * (attempt + 1))
                        continue
                    else:
                        raise PermissionError(
                            f"Could not save to '{self.file_path}'. Please ensure the file is not currently open in another program (e.g. Microsoft Excel)."
                        ) from pe
                finally:
                    if wb is not None:
                        try:
                            wb.close()
                        except Exception:
                            pass


if __name__ == "__main__":
    mgr = ExcelManager("./output/test_leads.xlsx")
    sample_lead = {
        "Name": "Rohan Mehta",
        "Phone no.": "",
        "Source": "Instagram",
        "Last Discussed topic": "Pricing for starter and growth plans and WhatsApp automation demo",
        "Status": "Interested / Follow-up Scheduled",
        "Assigned to": "Priya",
        "Follow up?": "WhatsApp follow up tomorrow evening with case studies"
    }
    mgr.append_lead(sample_lead)
    print("Appended test lead to ./output/test_leads.xlsx")
